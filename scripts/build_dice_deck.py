from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/workspaces/abc")
TEMPLATE = ROOT / "DICE Challenge S3  Template for Case studies submission  (Presentation).pptx"
OUTPUT = ROOT / "DICE Challenge S3 - Diced Cubers Final Submission.pptx"


PLUM = RGBColor(84, 12, 74)
PLUM_DARK = RGBColor(59, 7, 53)
ORANGE = RGBColor(245, 158, 11)
PINK = RGBColor(242, 74, 120)
ROSE = RGBColor(224, 85, 121)
CREAM = RGBColor(255, 247, 224)
INK = RGBColor(34, 22, 40)
MUTED = RGBColor(95, 78, 89)
OLIVE = RGBColor(194, 198, 79)
CYAN = RGBColor(130, 224, 242)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(255, 240, 214)
GRID = RGBColor(222, 206, 192)

FONT_HEAD = "Aptos Display"
FONT_BODY = "Aptos"


def remove_slide(prs: Presentation, index: int) -> None:
    slide_id_list = prs.slides._sldIdLst
    slides = list(slide_id_list)
    slide_id_list.remove(slides[index])


def set_bg(shape, color: RGBColor, transparency: float = 0.0) -> None:
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    fill.transparency = transparency
    shape.line.fill.background()


def add_box(slide, left, top, width, height, color, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, transparency=0.0):
    shape = slide.shapes.add_shape(radius, left, top, width, height)
    set_bg(shape, color, transparency)
    return shape


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=18,
    bold=False,
    color=INK,
    font=FONT_BODY,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.08,
    italic=False,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin * 0.7)
    tf.margin_bottom = Inches(margin * 0.4)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    items,
    *,
    size=14,
    color=INK,
    bullet_color=None,
    font=FONT_BODY,
    spacing=1.05,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.02)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(2)
        p.line_spacing = spacing
        run = p.add_run()
        bullet = "• "
        run.text = bullet + item
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if bullet_color:
            bullet_run = p.runs[0]
            bullet_run.font.color.rgb = bullet_color
    return box


def add_stat_chip(slide, left, top, width, title, body, color):
    add_box(slide, left, top, width, Inches(0.95), color)
    add_text(slide, left + Inches(0.08), top + Inches(0.08), width - Inches(0.16), Inches(0.28), title, size=11, bold=True, color=WHITE, font=FONT_BODY)
    add_text(slide, left + Inches(0.08), top + Inches(0.33), width - Inches(0.16), Inches(0.45), body, size=16, bold=True, color=WHITE, font=FONT_HEAD)


def add_label(slide, left, top, width, text, color=PLUM):
    pill = add_box(slide, left, top, width, Inches(0.35), color, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    add_text(slide, left + Inches(0.02), top + Inches(0.02), width - Inches(0.04), Inches(0.26), text, size=9, bold=True, color=WHITE, font=FONT_BODY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.0)
    return pill


def add_flow_card(slide, left, top, width, title, body, color):
    card = add_box(slide, left, top, width, Inches(1.05), WHITE, transparency=0.0)
    card.line.color.rgb = color
    card.line.width = Pt(1.5)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, Inches(0.24))
    set_bg(band, color)
    add_text(slide, left + Inches(0.06), top + Inches(0.05), width - Inches(0.12), Inches(0.20), title, size=11, bold=True, color=WHITE, font=FONT_BODY, margin=0.0)
    add_text(slide, left + Inches(0.06), top + Inches(0.30), width - Inches(0.12), Inches(0.68), body, size=10.5, color=INK, font=FONT_BODY)
    return card


def add_footer(slide, text):
    add_text(slide, Inches(0.55), Inches(10.64), Inches(18.0), Inches(0.28), text, size=8, color=MUTED, font=FONT_BODY)


def decorate_content_slide(slide, title, kicker=None):
    add_label(slide, Inches(0.62), Inches(0.58), Inches(1.95), "DICED CUBERS | IIT GUWAHATI")
    if kicker:
        add_label(slide, Inches(15.75), Inches(0.58), Inches(2.65), kicker, color=ORANGE)
    add_text(slide, Inches(0.62), Inches(1.1), Inches(11.9), Inches(0.8), title, size=25, bold=True, color=PLUM_DARK, font=FONT_HEAD)


def build_slide_1(slide):
    decorate_content_slide(slide, "Grow BPC Affiliate NMV by fixing first-purchase trust", "ROUND 1 | VALUE CHAIN")
    add_text(
        slide,
        Inches(0.68),
        Inches(1.78),
        Inches(9.4),
        Inches(0.62),
        "Research spine: market growth, competitor creator programs, and influencer-trust literature all point to the same gap: demand exists, but confidence breaks before checkout.",
        size=13,
        color=MUTED,
    )

    add_stat_chip(slide, Inches(0.72), Inches(2.35), Inches(2.75), "CATEGORY TAILWIND", "$39B India BPC by FY30", PLUM)
    add_stat_chip(slide, Inches(3.65), Inches(2.35), Inches(2.75), "DEMAND SHIFT", "50% YoY beauty growth; non-metros drive ~2/3 searches", PINK)
    add_stat_chip(slide, Inches(6.58), Inches(2.35), Inches(2.75), "COMPETITOR SIGNAL", "Nykaa + Myntra are scaling beauty creator ecosystems fast", ORANGE)

    panel = add_box(slide, Inches(13.15), Inches(1.72), Inches(5.5), Inches(3.2), PLUM_DARK)
    panel.fill.transparency = 0.03
    add_text(slide, Inches(13.4), Inches(1.95), Inches(5.0), Inches(0.4), "Why Meesho under-indexes in BPC affiliate today", size=17, bold=True, color=WHITE, font=FONT_HEAD)
    add_bullets(
        slide,
        Inches(13.35),
        Inches(2.42),
        Inches(5.05),
        Inches(2.2),
        [
            "Brands want credible demos, controlled sampling, and clean attribution before backing creators at scale.",
            "Activated creators face high trial cost and low confidence about what to post, whom to target, and how to monetize consistently.",
            "Users need proof on shade, skin or hair fit, authenticity, and value before a first beauty purchase.",
        ],
        size=12.5,
        color=WHITE,
    )

    add_text(slide, Inches(0.72), Inches(4.0), Inches(5.4), Inches(0.35), "End-to-end BPC influencer value chain", size=16, bold=True, color=PLUM_DARK, font=FONT_HEAD)
    stages = [
        ("Brand selects SKUs", "Hero products, claims, samples, target cohort.", ORANGE),
        ("Creator trials", "Receives samples; decides if product is content-worthy.", PINK),
        ("Proof content", "Demo, routine, before-after, price-value story.", PLUM),
        ("Traffic clickout", "IG reel, auto-DM, storefront, or link in bio.", CYAN),
        ("Trust-rich PDP", "Suitability answers and social proof before payment.", OLIVE),
        ("Repeat loop", "Satisfied buyer reorders and discovers adjacent SKUs.", ROSE),
    ]
    start = Inches(0.72)
    card_w = Inches(2.85)
    gap = Inches(0.18)
    y = Inches(4.35)
    for idx, (title, body, color) in enumerate(stages):
        x = start + idx * (card_w + gap)
        add_flow_card(slide, x, y, card_w, title, body, color)
        if idx < len(stages) - 1:
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                x + card_w,
                y + Inches(0.52),
                x + card_w + gap,
                y + Inches(0.52),
            )
            connector.line.color.rgb = PINK
            connector.line.width = Pt(1.5)

    add_box(slide, Inches(0.72), Inches(5.72), Inches(17.65), Inches(1.28), SOFT, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    add_text(slide, Inches(0.95), Inches(5.93), Inches(17.1), Inches(0.28), "Core insight", size=11, bold=True, color=PLUM, font=FONT_BODY)
    add_text(
        slide,
        Inches(0.95),
        Inches(6.16),
        Inches(17.1),
        Inches(0.62),
        "BPC does not fail on awareness; it fails when creators cannot cheaply generate credible proof and users cannot quickly judge suitability. Fix that single choke point, and creator retention plus NMV both move.",
        size=17,
        bold=True,
        color=INK,
        font=FONT_HEAD,
    )
    add_footer(
        slide,
        "Sources used in deck research: DICE case brief; 1Lattice via Economic Times (Mar 2026); Flipkart Beauty via Times of India (Jun 2026); Myntra creator program via ET; Nykaa beauty incubator via TOI; influencer trust literature.",
    )


def build_slide_2(slide):
    decorate_content_slide(slide, "Prioritize levers that improve all three sides at once", "ROUND 1 | PRIORITIZATION")
    add_text(
        slide,
        Inches(0.68),
        Inches(1.82),
        Inches(8.8),
        Inches(0.6),
        "North-star math: BPC affiliate NMV = active creators x quality posts x click-through x trust-led conversion x repeat rate.",
        size=14,
        color=MUTED,
    )

    matrix = add_box(slide, Inches(0.78), Inches(2.45), Inches(10.5), Inches(6.15), WHITE, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    matrix.line.color.rgb = GRID
    matrix.line.width = Pt(1.2)
    left = Inches(1.35)
    top = Inches(3.05)
    width = Inches(8.95)
    height = Inches(4.9)
    h_mid = top + height / 2
    v_mid = left + width / 2

    line1 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, h_mid, left + width, h_mid)
    line1.line.color.rgb = GRID
    line1.line.width = Pt(1.2)
    line2 = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, v_mid, top, v_mid, top + height)
    line2.line.color.rgb = GRID
    line2.line.width = Pt(1.2)

    add_text(slide, Inches(0.98), Inches(2.62), Inches(4.1), Inches(0.25), "3-sided impact", size=10, bold=True, color=MUTED, font=FONT_BODY)
    add_text(slide, Inches(9.2), Inches(8.0), Inches(1.5), Inches(0.25), "Ease / speed", size=10, bold=True, color=MUTED, font=FONT_BODY, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(1.05), Inches(3.2), Inches(1.3), Inches(0.3), "HIGH", size=10, bold=True, color=PLUM, font=FONT_BODY)
    add_text(slide, Inches(1.08), Inches(7.62), Inches(1.3), Inches(0.3), "LOW", size=10, bold=True, color=PLUM, font=FONT_BODY)
    add_text(slide, Inches(1.55), Inches(8.0), Inches(1.8), Inches(0.3), "SLOW", size=10, bold=True, color=PLUM, font=FONT_BODY)
    add_text(slide, Inches(8.55), Inches(8.0), Inches(1.8), Inches(0.3), "FAST", size=10, bold=True, color=PLUM, font=FONT_BODY)

    def bubble(x, y, w, text, color):
        shape = add_box(slide, Inches(x), Inches(y), Inches(w), Inches(0.55), color)
        shape.fill.transparency = 0.03
        add_text(slide, Inches(x) + Inches(0.05), Inches(y) + Inches(0.07), Inches(w) - Inches(0.1), Inches(0.36), text, size=10.5, bold=True, color=WHITE, font=FONT_BODY, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.0)

    bubble(6.35, 3.6, 3.15, "1. Creator Proof Kits", PINK)
    bubble(7.0, 4.75, 2.55, "2. PDP Suitability Layer", PLUM)
    bubble(5.2, 5.55, 2.95, "3. Creator-Brand Match Engine", ORANGE)
    bubble(3.0, 4.4, 2.65, "4. Tiered Incentive Redesign", ROSE)
    bubble(2.65, 6.15, 2.9, "5. Live Q&A / Consult", OLIVE)
    bubble(4.35, 6.85, 3.05, "6. Authenticity / Return Promise", CYAN)

    side = add_box(slide, Inches(11.65), Inches(2.45), Inches(6.95), Inches(6.15), PLUM_DARK, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    side.fill.transparency = 0.02
    add_text(slide, Inches(11.95), Inches(2.72), Inches(6.3), Inches(0.35), "Why Creator Proof Kits rank #1", size=18, bold=True, color=WHITE, font=FONT_HEAD)
    add_bullets(
        slide,
        Inches(11.92),
        Inches(3.17),
        Inches(6.15),
        Inches(1.9),
        [
            "Fastest path to unlock both activation and trust without waiting for a heavy-tech build.",
            "Creates content, creator habit, and brand proof together, instead of fixing only one side.",
            "Generates structured data Meesho can later use for ranking, matching, and personalization.",
        ],
        size=13,
        color=WHITE,
    )

    add_text(slide, Inches(11.95), Inches(5.28), Inches(6.0), Inches(0.28), "Decision rule", size=11, bold=True, color=CYAN, font=FONT_BODY)
    add_text(
        slide,
        Inches(11.95),
        Inches(5.55),
        Inches(6.0),
        Inches(0.88),
        "Prioritize levers that move at least three NMV drivers at once. Pure incentive changes may spike posting, but trust-led conversion still remains broken.",
        size=14,
        color=WHITE,
        font=FONT_BODY,
    )

    add_box(slide, Inches(11.95), Inches(6.65), Inches(6.0), Inches(1.38), SOFT, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    add_text(slide, Inches(12.18), Inches(6.87), Inches(5.5), Inches(0.23), "Prioritized stack", size=11, bold=True, color=PLUM, font=FONT_BODY)
    add_bullets(
        slide,
        Inches(12.15),
        Inches(7.1),
        Inches(5.5),
        Inches(0.78),
        [
            "30 days: Creator Proof Kits",
            "60 days: suitability snippets on PDP + storefronts",
            "90 days: creator-brand matching + smarter budget allocation",
        ],
        size=12,
        color=INK,
    )
    add_footer(
        slide,
        "Framework used: impact on brands + creators + users, expected NMV lift, implementation dependency, and speed to pilot within a 30-day window.",
    )


def build_slide_3(slide):
    decorate_content_slide(slide, "Quick win: launch Meesho Beauty Proof Kits in 30 days", "ROUND 1 | SOLUTION DESIGN")
    add_box(slide, Inches(0.72), Inches(1.9), Inches(18.0), Inches(0.78), PLUM, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    add_text(
        slide,
        Inches(0.95),
        Inches(2.1),
        Inches(17.4),
        Inches(0.35),
        "A curated starter program where creators receive hero BPC samples, post a structured proof format, and send traffic to trust-rich product pages with clean attribution.",
        size=15,
        color=WHITE,
        bold=True,
        font=FONT_BODY,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    cards = [
        ("1. Sample Wallet", "Brands co-fund 20 hero SKUs in minis or low-cost trials. Meesho only pushes products with repeat potential and clean margins.", ORANGE),
        ("2. Proof Template", "Every creator post captures concern, skin or hair type, shade or finish, price-value cue, and a short disclosure. Content becomes comparable.", PINK),
        ("3. Trust Layer", "On storefront / PDP, users see creator demo clip, fit tags, FAQs, and 'best for' guidance before checkout.", PLUM),
        ("4. Incentive Engine", "Starter bounty for first 25 orders, retention bonus at day 30, and brand dashboard that shows real sales not vanity views.", OLIVE),
    ]
    y = Inches(3.0)
    for idx, (title, body, color) in enumerate(cards):
        top = y + idx * Inches(1.28)
        card = add_box(slide, Inches(0.82), top, Inches(8.7), Inches(1.0), WHITE, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
        card.line.color.rgb = color
        card.line.width = Pt(1.4)
        band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.82), top, Inches(2.05), Inches(1.0))
        set_bg(band, color)
        add_text(slide, Inches(0.97), top + Inches(0.18), Inches(1.7), Inches(0.45), title, size=13, bold=True, color=WHITE, font=FONT_HEAD, margin=0.0)
        add_text(slide, Inches(3.02), top + Inches(0.16), Inches(6.15), Inches(0.56), body, size=12.3, color=INK, font=FONT_BODY)

    impact = add_box(slide, Inches(10.0), Inches(3.0), Inches(8.6), Inches(3.5), PLUM_DARK, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    impact.fill.transparency = 0.02
    add_text(slide, Inches(10.28), Inches(3.24), Inches(8.0), Inches(0.35), "Pilot design and expected impact", size=18, bold=True, color=WHITE, font=FONT_HEAD)
    add_bullets(
        slide,
        Inches(10.25),
        Inches(3.72),
        Inches(7.85),
        Inches(1.25),
        [
            "Pilot cohort: 10 brands, 20 hero SKUs, 200 activated creators, 4 weeks.",
            "Success metric: uplift on creator-led BPC NMV versus similar untreated SKUs and creator cohorts.",
            "Hypothesis: +12-18% BPC affiliate NMV, +15-20% 30-day creator retention, +8-12% PDP-to-order conversion on pilot SKUs.",
        ],
        size=12.7,
        color=WHITE,
    )
    add_box(slide, Inches(10.28), Inches(5.15), Inches(7.95), Inches(1.02), SOFT, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    add_text(slide, Inches(10.5), Inches(5.34), Inches(7.45), Inches(0.25), "Why this wins the round", size=11, bold=True, color=PLUM, font=FONT_BODY)
    add_text(
        slide,
        Inches(10.5),
        Inches(5.56),
        Inches(7.3),
        Inches(0.42),
        "It is creator-native, low-tech enough to ship fast, and strong enough to become Meesho's long-term beauty content moat.",
        size=15,
        bold=True,
        color=INK,
        font=FONT_HEAD,
    )

    add_text(slide, Inches(10.0), Inches(6.8), Inches(4.0), Inches(0.3), "30-day rollout", size=16, bold=True, color=PLUM_DARK, font=FONT_HEAD)
    timeline_y = Inches(7.25)
    xs = [Inches(10.1), Inches(12.95), Inches(15.8)]
    titles = ["Week 1", "Week 2-3", "Week 4"]
    bodies = [
        "Select brands, lock hero SKUs, issue sample wallet.",
        "Creators publish proof-form videos; Meesho tags fit metadata.",
        "Rank by CTR/CVR, fund winners harder, and codify repeat playbook.",
    ]
    colors = [ORANGE, PINK, CYAN]
    for x, title, body, color in zip(xs, titles, bodies, colors):
        card = add_box(slide, x, timeline_y, Inches(2.45), Inches(1.1), WHITE, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
        card.line.color.rgb = color
        card.line.width = Pt(1.3)
        add_text(slide, x + Inches(0.08), timeline_y + Inches(0.14), Inches(2.2), Inches(0.22), title, size=11, bold=True, color=color, font=FONT_BODY)
        add_text(slide, x + Inches(0.08), timeline_y + Inches(0.4), Inches(2.22), Inches(0.55), body, size=10.6, color=INK, font=FONT_BODY)

    add_box(slide, Inches(0.82), Inches(8.65), Inches(17.78), Inches(0.82), WHITE, radius=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    add_text(slide, Inches(1.05), Inches(8.89), Inches(17.2), Inches(0.25), "Team", size=11, bold=True, color=PLUM, font=FONT_BODY)
    add_text(
        slide,
        Inches(1.05),
        Inches(9.09),
        Inches(17.2),
        Inches(0.28),
        "Samiksha Mitra | Diced Cubers | IIT Guwahati | Proposed thesis: win beauty commerce by turning creator proof into a repeatable acquisition engine.",
        size=13,
        color=INK,
        font=FONT_BODY,
    )
    add_footer(slide, "Impact figures are pilot hypotheses derived from the case's stated bottlenecks and the NMV equation above; they are intended for test-and-learn validation in Round 2.")


def main():
    prs = Presentation(str(TEMPLATE))

    for idx in sorted([4, 0], reverse=True):
        remove_slide(prs, idx)

    slides = list(prs.slides)
    build_slide_1(slides[0])
    build_slide_2(slides[1])
    build_slide_3(slides[2])
    prs.save(str(OUTPUT))


if __name__ == "__main__":
    main()
